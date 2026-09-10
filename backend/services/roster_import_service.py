from __future__ import annotations

import csv
from dataclasses import dataclass, field
from hashlib import sha256
from io import BytesIO, StringIO
from pathlib import PurePosixPath
import re
from typing import Any, Iterable
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

from core.config import supabase_db as supabase


MAX_ROSTER_FILE_BYTES = 5 * 1024 * 1024
MAX_ROSTER_ROWS = 1000
MAX_SOURCE_ROWS = 5000
MAX_SOURCE_COLUMNS = 100
MAX_CELL_CHARACTERS = 500
MAX_XLSX_ENTRIES = 2000
MAX_XLSX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_XLSX_ENTRY_BYTES = 20 * 1024 * 1024
MAX_XLSX_COMPRESSION_RATIO = 1000

CSV_CONTENT_TYPES = {
    "text/csv",
    "application/csv",
    "text/plain",
    "application/vnd.ms-excel",
    "application/octet-stream",
}
XLSX_CONTENT_TYPES = {
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/zip",
    "application/octet-stream",
}


class RosterImportError(ValueError):
    def __init__(self, message: str, status_code: int = 400):
        super().__init__(message)
        self.message = message
        self.status_code = status_code


@dataclass
class ParsedRosterRow:
    source_row: int
    student_id: str
    display_student_id: str
    full_name: str
    email: str
    academic_year: int | None
    class_level: str | None
    errors: list[str] = field(default_factory=list)
    action: str | None = None

    def as_dict(self) -> dict[str, Any]:
        return {
            "source_row": self.source_row,
            "student_id": self.student_id,
            "display_student_id": self.display_student_id,
            "full_name": self.full_name,
            "email": self.email,
            "academic_year": self.academic_year,
            "class_level": self.class_level,
            "errors": self.errors,
            "action": self.action,
        }

    def commit_dict(self) -> dict[str, Any]:
        return {
            "source_row": self.source_row,
            "student_id": self.student_id,
            "full_name": self.full_name,
            "email": self.email,
            "academic_year": self.academic_year,
            "class_level": self.class_level,
        }


@dataclass
class ParsedRoster:
    filename: str
    file_type: str
    digest: str
    sheet_name: str | None
    encoding: str | None
    header_row: int
    detected_course_code: str | None
    detected_course_name: str | None
    detected_academic_year: int | None
    rows: list[ParsedRosterRow]
    warnings: list[str]


def _clean_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).replace("\u200b", "").replace("\ufeff", "")
    if len(text) > MAX_CELL_CHARACTERS:
        raise RosterImportError("พบข้อมูลในช่องที่ยาวผิดปกติ ไฟล์อาจไม่ใช่ใบรายชื่อ")
    return re.sub(r"\s+", " ", text).strip()


def _header_key(value: Any) -> str:
    return re.sub(r"[\s_\-–—./]+", "", _clean_text(value).casefold())


HEADER_ALIASES = {
    "student_id": {
        "studentid", "รหัสนักศึกษา", "เลขประจำตัวนักศึกษา",
    },
    "full_name": {
        "fullname", "ชื่อสกุล", "ชื่อนามสกุล", "ชื่อและนามสกุล",
    },
    "email": {"email", "emailaddress", "อีเมล", "อีเมลนักศึกษา"},
    "academic_year": {"academicyear", "ชั้นปี"},
    "class_level": {"classlevel", "กลุ่มเรียน", "ห้องเรียน", "กลุ่ม"},
}


def _find_header(rows: list[list[Any]]) -> tuple[int, int, dict[str, int]]:
    for index, row in enumerate(rows[:50]):
        next_row = rows[index + 1] if index + 1 < len(rows) else []
        mapping: dict[str, int] = {}
        used_second_header = False
        for column in range(min(MAX_SOURCE_COLUMNS, max(len(row), len(next_row)))):
            current = _header_key(row[column] if column < len(row) else None)
            combined = _header_key(
                f"{row[column] if column < len(row) else ''} "
                f"{next_row[column] if column < len(next_row) else ''}"
            )
            for field_name, aliases in HEADER_ALIASES.items():
                if field_name in mapping:
                    continue
                if current in aliases:
                    mapping[field_name] = column
                elif combined in aliases:
                    mapping[field_name] = column
                    used_second_header = True
        if "student_id" in mapping and "full_name" in mapping:
            return index, index + (2 if used_second_header else 1), mapping
    raise RosterImportError(
        "ไม่พบหัวตารางรายชื่อ ต้องมีคอลัมน์รหัสนักศึกษาและชื่อ-นามสกุล"
    )


def _extract_metadata(rows: list[list[Any]]) -> tuple[str | None, str | None, int | None]:
    preamble = "\n".join(
        " ".join(_clean_text(cell) for cell in row if _clean_text(cell))
        for row in rows[:20]
    )
    year_match = re.search(r"ภาคการศึกษา\s*\d+\s*/\s*(25\d{2}|20\d{2})", preamble)
    academic_year = int(year_match.group(1)) if year_match else None

    course_match = re.search(
        r"วิชา\s+([0-9]{6,12})\s+(.+?)(?:\s+\d+\s*\(|\s+ตอน\s+|$)",
        preamble,
    )
    if not course_match:
        return None, None, academic_year
    return course_match.group(1), _clean_text(course_match.group(2)), academic_year


def _derive_study_year(student_id: str, academic_year: int | None) -> int | None:
    if academic_year is None or len(student_id) != 13:
        return None
    admission_year = int(student_id[:2])
    current_short_year = academic_year % 100
    derived = (current_short_year - admission_year) % 100 + 1
    return derived if 1 <= derived <= 8 else None


def _parse_academic_year(value: Any) -> int | None:
    text = _clean_text(value)
    if not text:
        return None
    match = re.fullmatch(r"(?:ปี\s*)?([1-8])", text)
    return int(match.group(1)) if match else -1


def _formula_like(value: str) -> bool:
    return bool(re.match(r"^[\s\t\r]*[=+\-@]", value))


def _rows_to_roster(
    rows: list[list[Any]],
    *,
    filename: str,
    file_type: str,
    digest: str,
    sheet_name: str | None,
    encoding: str | None,
) -> ParsedRoster:
    header_index, data_start, columns = _find_header(rows)
    course_code, course_name, source_academic_year = _extract_metadata(rows)
    parsed: list[ParsedRosterRow] = []

    for index, source in enumerate(rows[data_start:], start=data_start + 1):
        if any(_clean_text(value).startswith("พิมพ์ ณ วันที่") for value in source):
            break

        def value_for(field_name: str) -> Any:
            column = columns.get(field_name)
            return source[column] if column is not None and column < len(source) else None

        raw_student_id = _clean_text(value_for("student_id"))
        full_name = _clean_text(value_for("full_name"))
        class_level = _clean_text(value_for("class_level")) or None
        supplied_email = _clean_text(value_for("email")).casefold()
        academic_year = _parse_academic_year(value_for("academic_year"))

        if not any((raw_student_id, full_name, class_level, supplied_email)):
            continue

        student_id = re.sub(r"\D", "", raw_student_id)
        email = supplied_email or (
            f"s{student_id}@email.kmutnb.ac.th" if len(student_id) == 13 else ""
        )
        if academic_year is None:
            academic_year = _derive_study_year(student_id, source_academic_year)

        errors: list[str] = []
        if not re.fullmatch(r"[0-9]{13}", student_id):
            errors.append("รหัสนักศึกษาต้องมีตัวเลข 13 หลัก")
        if not 2 <= len(full_name) <= 150:
            errors.append("ชื่อ-นามสกุลต้องมีความยาว 2–150 ตัวอักษร")
        elif _formula_like(full_name):
            errors.append("ชื่อ-นามสกุลมีอักขระขึ้นต้นที่ไม่ปลอดภัย")
        if email and not re.fullmatch(r"[^@]+@email\.kmutnb\.ac\.th", email):
            errors.append("อีเมลต้องอยู่ในโดเมนนักศึกษา KMUTNB")
        if academic_year == -1:
            errors.append("ชั้นปีต้องเป็นตัวเลข 1–8")
            academic_year = None
        if class_level and len(class_level) > 50:
            errors.append("กลุ่มเรียนต้องไม่เกิน 50 ตัวอักษร")
        elif class_level and _formula_like(class_level):
            errors.append("กลุ่มเรียนมีอักขระขึ้นต้นที่ไม่ปลอดภัย")

        parsed.append(ParsedRosterRow(
            source_row=index,
            student_id=student_id,
            display_student_id=raw_student_id or student_id,
            full_name=full_name,
            email=email,
            academic_year=academic_year,
            class_level=class_level,
            errors=errors,
        ))
        if len(parsed) > MAX_ROSTER_ROWS:
            raise RosterImportError(
                f"รองรับรายชื่อไม่เกิน {MAX_ROSTER_ROWS:,} คนต่อไฟล์",
                status_code=413,
            )

    if not parsed:
        raise RosterImportError("ไม่พบแถวรายชื่อนักศึกษาในไฟล์")

    duplicate_ids: dict[str, list[ParsedRosterRow]] = {}
    for row in parsed:
        if re.fullmatch(r"[0-9]{13}", row.student_id):
            duplicate_ids.setdefault(row.student_id, []).append(row)
    for duplicates in duplicate_ids.values():
        if len(duplicates) > 1:
            for row in duplicates:
                row.errors.append("รหัสนักศึกษาซ้ำในไฟล์")

    warnings: list[str] = []
    if course_code is None:
        warnings.append("ไฟล์ไม่มีข้อมูลรหัสวิชา ระบบจะใช้รายวิชาที่เลือกอยู่")
    if "email" not in columns:
        warnings.append("สร้างอีเมลนักศึกษาจากรหัสในรูปแบบ s<รหัส>@email.kmutnb.ac.th")
    if source_academic_year is None and "academic_year" not in columns:
        warnings.append("ไม่พบปีการศึกษา จึงยังไม่กำหนดชั้นปีให้อัตโนมัติ")

    return ParsedRoster(
        filename=filename,
        file_type=file_type,
        digest=digest,
        sheet_name=sheet_name,
        encoding=encoding,
        header_row=header_index + 1,
        detected_course_code=course_code,
        detected_course_name=course_name,
        detected_academic_year=source_academic_year,
        rows=parsed,
        warnings=warnings,
    )


def _validate_xlsx_archive(content: bytes) -> None:
    try:
        with ZipFile(BytesIO(content)) as archive:
            infos = archive.infolist()
            if not infos or len(infos) > MAX_XLSX_ENTRIES:
                raise RosterImportError("โครงสร้างไฟล์ XLSX มีจำนวนส่วนประกอบผิดปกติ")
            names = {info.filename for info in infos}
            if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
                raise RosterImportError("ไฟล์ไม่ใช่ XLSX ที่สมบูรณ์")

            total_uncompressed = 0
            for info in infos:
                path = PurePosixPath(info.filename)
                if path.is_absolute() or ".." in path.parts or info.flag_bits & 0x1:
                    raise RosterImportError("ไฟล์ XLSX มีโครงสร้างที่ไม่ปลอดภัย")
                if info.filename.startswith("xl/externalLinks/") or info.filename.endswith("vbaProject.bin"):
                    raise RosterImportError("ไฟล์ XLSX ต้องไม่มีลิงก์ภายนอกหรือ Macro")
                if info.file_size > MAX_XLSX_ENTRY_BYTES:
                    raise RosterImportError("ส่วนประกอบภายใน XLSX มีขนาดใหญ่เกินกำหนด", 413)
                total_uncompressed += info.file_size
                if info.file_size and info.file_size / max(1, info.compress_size) > MAX_XLSX_COMPRESSION_RATIO:
                    raise RosterImportError("ไฟล์ XLSX มีอัตราการบีบอัดผิดปกติ", 413)
            if total_uncompressed > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise RosterImportError("ข้อมูลภายใน XLSX มีขนาดใหญ่เกินกำหนด", 413)
    except BadZipFile as exc:
        raise RosterImportError("ไฟล์ XLSX เสียหายหรือไม่ใช่ไฟล์ Excel") from exc


def _read_xlsx_rows(content: bytes) -> tuple[list[list[Any]], str]:
    _validate_xlsx_archive(content)
    try:
        workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except Exception as exc:
        raise RosterImportError("ไม่สามารถอ่านไฟล์ XLSX ได้") from exc

    try:
        if len(workbook.worksheets) > 10:
            raise RosterImportError("ไฟล์ XLSX มีจำนวนชีตมากเกินกำหนด")
        for worksheet in workbook.worksheets:
            if worksheet.max_row > MAX_SOURCE_ROWS or worksheet.max_column > MAX_SOURCE_COLUMNS:
                raise RosterImportError("ชีตมีจำนวนแถวหรือคอลัมน์มากเกินกำหนด", 413)
            rows: list[list[Any]] = []
            for cells in worksheet.iter_rows(
                min_row=1,
                max_row=min(worksheet.max_row, MAX_SOURCE_ROWS),
                max_col=min(worksheet.max_column, MAX_SOURCE_COLUMNS),
            ):
                row: list[Any] = []
                for cell in cells:
                    if cell.data_type == "f":
                        raise RosterImportError("ไฟล์รายชื่อต้องไม่มีสูตรคำนวณ")
                    row.append(cell.value)
                rows.append(row)
            try:
                _find_header(rows)
                return rows, worksheet.title
            except RosterImportError:
                continue
    finally:
        workbook.close()
    raise RosterImportError("ไม่พบชีตที่มีหัวตารางรายชื่อนักศึกษา")


def _read_csv_rows(content: bytes) -> tuple[list[list[str]], str]:
    text = None
    encoding = None
    for candidate in ("utf-8-sig", "utf-8", "cp874"):
        try:
            text = content.decode(candidate)
            encoding = candidate
            break
        except UnicodeDecodeError:
            continue
    if text is None or encoding is None:
        raise RosterImportError("ไฟล์ CSV ต้องเข้ารหัสเป็น UTF-8 หรือ Windows Thai (CP874)")
    if "\x00" in text:
        raise RosterImportError("ไฟล์ CSV มีข้อมูลไบนารีที่ไม่ปลอดภัย")

    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel

    previous_limit = csv.field_size_limit()
    csv.field_size_limit(10_000)
    try:
        rows: list[list[str]] = []
        reader = csv.reader(StringIO(text, newline=""), dialect=dialect, strict=True)
        for row in reader:
            if len(row) > MAX_SOURCE_COLUMNS:
                raise RosterImportError("CSV มีจำนวนคอลัมน์มากเกินกำหนด", 413)
            if any(len(cell) > MAX_CELL_CHARACTERS for cell in row):
                raise RosterImportError("CSV มีข้อมูลในช่องที่ยาวผิดปกติ", 413)
            rows.append(row)
            if len(rows) > MAX_SOURCE_ROWS:
                raise RosterImportError("CSV มีจำนวนแถวมากเกินกำหนด", 413)
    except csv.Error as exc:
        raise RosterImportError("โครงสร้าง CSV ไม่ถูกต้อง") from exc
    finally:
        csv.field_size_limit(previous_limit)
    return rows, encoding


def parse_roster_file(content: bytes, filename: str, content_type: str | None) -> ParsedRoster:
    if not content:
        raise RosterImportError("ไฟล์ว่างเปล่า")
    if len(content) > MAX_ROSTER_FILE_BYTES:
        raise RosterImportError("ไฟล์รายชื่อต้องมีขนาดไม่เกิน 5 MB", 413)

    safe_name = PurePosixPath(filename.replace("\\", "/")).name
    suffix = PurePosixPath(safe_name).suffix.casefold()
    normalized_type = (content_type or "application/octet-stream").split(";", 1)[0].casefold()
    digest = sha256(content).hexdigest()

    if suffix == ".xlsx":
        if normalized_type not in XLSX_CONTENT_TYPES or not content.startswith(b"PK"):
            raise RosterImportError("ชนิดข้อมูลภายในไฟล์ไม่ตรงกับนามสกุล XLSX")
        rows, sheet_name = _read_xlsx_rows(content)
        return _rows_to_roster(
            rows,
            filename=safe_name,
            file_type="xlsx",
            digest=digest,
            sheet_name=sheet_name,
            encoding=None,
        )

    if suffix == ".csv":
        if normalized_type not in CSV_CONTENT_TYPES or content.startswith(b"PK"):
            raise RosterImportError("ชนิดข้อมูลภายในไฟล์ไม่ตรงกับนามสกุล CSV")
        rows, encoding = _read_csv_rows(content)
        return _rows_to_roster(
            rows,
            filename=safe_name,
            file_type="csv",
            digest=digest,
            sheet_name=None,
            encoding=encoding,
        )

    raise RosterImportError("รองรับเฉพาะไฟล์ .xlsx และ .csv")


def _chunks(values: list[str], size: int = 100) -> Iterable[list[str]]:
    for index in range(0, len(values), size):
        yield values[index:index + size]


def _load_by_values(table: str, select: str, column: str, values: list[str]) -> list[dict]:
    result: list[dict] = []
    for chunk in _chunks(sorted(set(value for value in values if value))):
        response = supabase.table(table).select(select).in_(column, chunk).execute()
        result.extend(response.data or [])
    return result


def resolve_roster_actions(parsed: ParsedRoster, course_id: str) -> dict[str, int]:
    candidate_rows = [row for row in parsed.rows if not row.errors]
    student_ids = [row.student_id for row in candidate_rows]
    emails = [row.email for row in candidate_rows]

    profiles = _load_by_values(
        "profiles", "id, student_id, email, role", "student_id", student_ids
    )
    profiles.extend(_load_by_values(
        "profiles", "id, student_id, email, role", "email", emails
    ))
    profiles_by_student_id = {
        row.get("student_id"): row for row in profiles if row.get("student_id")
    }
    profiles_by_email = {
        str(row.get("email") or "").casefold(): row for row in profiles if row.get("email")
    }

    invites = _load_by_values(
        "profile_invites",
        "id, student_id, email, role, claimed_at",
        "student_id",
        student_ids,
    )
    invites.extend(_load_by_values(
        "profile_invites",
        "id, student_id, email, role, claimed_at",
        "email",
        emails,
    ))
    invites_by_student_id = {
        row.get("student_id"): row for row in invites if row.get("student_id")
    }
    invites_by_email = {
        str(row.get("email") or "").casefold(): row for row in invites if row.get("email")
    }

    profile_ids = sorted({str(row["id"]) for row in profiles if row.get("id")})
    enrolled_profile_ids: set[str] = set()
    for chunk in _chunks(profile_ids):
        response = (
            supabase.table("enrollments")
            .select("student_id")
            .eq("course_id", course_id)
            .in_("student_id", chunk)
            .execute()
        )
        enrolled_profile_ids.update(str(row["student_id"]) for row in (response.data or []))

    invite_ids = sorted({str(row["id"]) for row in invites if row.get("id")})
    linked_invite_ids: set[str] = set()
    for chunk in _chunks(invite_ids):
        response = (
            supabase.table("profile_invite_courses")
            .select("invite_id")
            .eq("course_id", course_id)
            .in_("invite_id", chunk)
            .execute()
        )
        linked_invite_ids.update(str(row["invite_id"]) for row in (response.data or []))

    summary = {
        "enroll": 0,
        "create_invite": 0,
        "link_invite": 0,
        "unchanged": 0,
        "invalid": 0,
    }
    for row in parsed.rows:
        if row.errors:
            row.action = "invalid"
            summary["invalid"] += 1
            continue

        profile = profiles_by_student_id.get(row.student_id)
        profile_email_match = profiles_by_email.get(row.email.casefold())
        if profile:
            if profile.get("role") != "student" or (
                profile_email_match and profile_email_match.get("id") != profile.get("id")
            ):
                row.errors.append("รหัสหรืออีเมลตรงกับบัญชีอื่นในระบบ")
                row.action = "invalid"
                summary["invalid"] += 1
                continue
            if str(profile["id"]) in enrolled_profile_ids:
                row.action = "unchanged"
                summary["unchanged"] += 1
            else:
                row.action = "enroll"
                summary["enroll"] += 1
            continue

        if profile_email_match:
            row.errors.append("อีเมลตรงกับบัญชีที่ใช้รหัสนักศึกษาอื่น")
            row.action = "invalid"
            summary["invalid"] += 1
            continue

        invite = invites_by_student_id.get(row.student_id)
        invite_email_match = invites_by_email.get(row.email.casefold())
        if invite:
            if invite.get("role") != "student" or invite.get("claimed_at") or (
                invite_email_match and invite_email_match.get("id") != invite.get("id")
            ):
                row.errors.append("รหัสหรืออีเมลตรงกับคำเชิญอื่นในระบบ")
                row.action = "invalid"
                summary["invalid"] += 1
                continue
            if str(invite["id"]) in linked_invite_ids:
                row.action = "unchanged"
                summary["unchanged"] += 1
            else:
                row.action = "link_invite"
                summary["link_invite"] += 1
            continue

        if invite_email_match:
            row.errors.append("อีเมลตรงกับคำเชิญที่ใช้รหัสนักศึกษาอื่น")
            row.action = "invalid"
            summary["invalid"] += 1
            continue

        row.action = "create_invite"
        summary["create_invite"] += 1

    return summary
